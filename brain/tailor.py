"""Tailored CV / resume PDFs per job (Goal 5), with two engines you can compare:

  career-ops       LLM fills career-ops' HTML CV template using the evaluation report + JD (anti-fabrication rules),
                   then Playwright renders the PDF.
  resume-matcher   Resume-Matcher's backend (started privately on its own port, settings from brain/.env only).

make_resume(job_id, engine) -> {"ok", "engine", "pdf_path", "error", "model"}   (engine "both" adds "results")
Files: <profile>/workspace/output/resumes/<job_id>/<engine>.pdf  (older versions kept in history/)
"""
import hashlib
import json
import os
import re
import shutil
import socket
import subprocess
import time
from pathlib import Path

import requests

from brain import config, jobs_db, nodeenv, profiles, sync_env
from brain.llm import LLMUnavailable
from brain.config import ROOT_DIR

CAREER_OPS = ROOT_DIR / "career-ops"
RM_BACKEND = ROOT_DIR / "Resume-Matcher" / "apps" / "backend"
ENGINES = ("career-ops", "resume-matcher")
RM_PORT = 8765
RM_FRONT_PORT = 8766
US_LIKE = {"us", "usa", "united states", "united states of america", "canada", "ca"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class TailorError(RuntimeError):
    pass


def out_dir(job_id: str, name: str | None = None) -> Path:
    d = profiles.workspace_path(name, "output") / "resumes" / job_id
    d.mkdir(parents=True, exist_ok=True)
    return d


def _archive(path: Path) -> None:
    if path.exists():
        h = path.parent / "history"
        h.mkdir(exist_ok=True)
        shutil.move(str(path), str(h / f"{path.stem}-{time.strftime('%Y%m%d-%H%M%S')}-{int(time.time() * 1000) % 1000:03d}{path.suffix}"))


def list_resumes(job_id: str, name: str | None = None) -> dict:
    d = out_dir(job_id, name)
    return {e: str(d / f"{e}.pdf") for e in ENGINES if (d / f"{e}.pdf").exists()}


def _job(job_id: str, name: str) -> dict:
    job = jobs_db.get_job(job_id, name)
    if not job:
        raise TailorError(f"No job '{job_id}'")
    if len(job.get("description") or "") < 700:                   # thin teaser -> full posting
        from brain import rank
        full = rank._fulltext(job["url"])
        if full:
            jobs_db.update_fields(job_id, name, description=full)
            job = {**job, "description": full}
    if len((job.get("description") or "").strip()) < 100:
        raise TailorError("This job has no usable description to tailor against. Open the posting and paste the text, or re-run discovery.")
    return job


def paper_format(name: str) -> str:
    country = str(profiles.load_profile(name)["prefs"].get("country", "")).strip().lower()
    return "letter" if country in US_LIKE else "a4"


# ================================================================== engine A: career-ops
def validate_html(html: str, cv_text: str = "") -> str | None:
    """Return a problem description, or None if the HTML looks like a finished CV."""
    h = html.strip()
    if not re.match(r"(?is)^\s*(<!doctype html|<html)", h):
        return "does not start with an HTML document"
    if "</html>" not in h.lower():
        return "HTML is cut off (no closing </html>)"
    left = re.findall(r"\{\{[A-Z0-9_]+\}\}", h)
    if left:
        return f"template placeholders left unfilled: {left[:3]}"
    if any(m in h for m in ("<function", "<parameter", "<tool_call")):
        return "contains tool-call text"
    if len(re.sub(r"<[^>]+>", " ", h)) < 800:
        return "almost no text"
    return None


def _run_node(args: list[str], timeout: int = 300) -> subprocess.CompletedProcess:
    return subprocess.run([nodeenv.node_bin(), *args], cwd=CAREER_OPS, env=nodeenv.env_with_node(),
                          capture_output=True, text=True, timeout=timeout)


def make_career_ops(job_id: str, name: str | None = None, models: list[str] | None = None, timeout: int = 600) -> dict:
    name = name or profiles.get_active()
    res = {"ok": False, "engine": "career-ops", "pdf_path": None, "error": None, "model": None}
    try:
        job = _job(job_id, name)
        report = job.get("report_path")
        if not report or not Path(report).exists():
            raise TailorError("This job has not been evaluated yet - run `hire.py rank` first (the tailoring uses its evaluation report).")
        if not nodeenv.node_bin():
            raise TailorError("Node.js >= 22 not found (nvm install 22).")
        profiles.link_active(name)
        cv_text = profiles.read_file(name, "cv.md")
        jd_file = CAREER_OPS / "jds" / f"{job_id}.txt"
        jd_file.parent.mkdir(exist_ok=True)
        jd_file.write_text(f"{job.get('title') or ''} at {job.get('company') or ''}\n\n{job['description']}", encoding="utf-8")
        outdir = CAREER_OPS / "output"
        before = {p: p.stat().st_mtime for p in outdir.glob("cv-*.html")} if outdir.exists() else {}
        found = {}

        def check(proc, model):
            cand = [p for p in outdir.glob("cv-*.html") if p.stat().st_mtime > before.get(p, 0)]
            if not cand:
                return "no HTML written"
            p = max(cand, key=lambda x: x.stat().st_mtime)
            bad = validate_html(p.read_text(encoding="utf-8", errors="replace"), cv_text)
            if bad:
                p.unlink(missing_ok=True)
                return f"unusable: {bad}"
            found["html"] = p
            return None

        try:
            _, model = sync_env.run_with_fallback([nodeenv.node_bin(), "openai-tailor.mjs", "--jd", str(jd_file), "--report", str(report)],
                                                  str(CAREER_OPS), timeout=timeout, models=models, validate=check)
        except LLMUnavailable as e:
            raise TailorError("No model produced a usable CV: " + "; ".join(f"{k}: {v}" for k, v in list(e.errors.items())[:3]))
        d = out_dir(job_id, name)
        pdf, html = d / "career-ops.pdf", d / "career-ops.html"
        from brain import reports
        num = reports.report_number(report)
        tmp_pdf = d / ".career-ops.tmp.pdf"
        args = ["generate-pdf.mjs", str(found["html"]), str(tmp_pdf), f"--format={paper_format(name)}"] + ([f"--report={num:03d}"] if num else [])
        p = _run_node(args)
        if p.returncode != 0 or not tmp_pdf.exists():
            tail = (p.stdout + p.stderr).strip().splitlines()[-4:]
            raise TailorError("PDF rendering failed: " + " | ".join(tail)[:300])
        _archive(pdf)
        _archive(html)
        tmp_pdf.replace(pdf)
        shutil.copy2(found["html"], html)
        res.update(ok=True, pdf_path=str(pdf), model=model, html_path=str(html))
    except (TailorError, subprocess.TimeoutExpired) as e:
        res["error"] = str(e) if isinstance(e, TailorError) else "timed out"
    return res


# ================================================================== engine B: Resume-Matcher
def cv_to_docx(cv_md: str, path: Path) -> Path:
    """Resume-Matcher only accepts PDF/DOCX uploads, so render cv.md as a plain DOCX."""
    import docx
    d = docx.Document()
    for raw in cv_md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            d.add_heading(re.sub(r"[*_`]", "", m.group(2)), level=min(len(m.group(1)), 3))
        elif re.match(r"^\s*[-*+]\s+", line):
            d.add_paragraph(re.sub(r"^\s*[-*+]\s+", "", re.sub(r"\*\*|__|`", "", line)), style="List Bullet")
        else:
            d.add_paragraph(re.sub(r"\*\*|__|`", "", line))
    d.save(str(path))
    return path


def _port_free(port: int) -> bool:
    with socket.socket() as s:
        return s.connect_ex(("127.0.0.1", port)) != 0


class Backend:
    """A private Resume-Matcher backend: started with LLM settings from brain/.env, stopped afterwards."""

    def __init__(self, model: str, port: int = RM_PORT):
        self.model, self.port, self.proc, self.front = model, port, None, None
        self.front_port = port + 1
        self.base = f"http://127.0.0.1:{port}/api/v1"

    def start(self, wait: int = 60) -> "Backend":
        if not (RM_BACKEND / ".venv" / "bin" / "python").exists():
            raise TailorError("Resume-Matcher is not installed (apps/backend/.venv missing). See plan.md Phase 0.")
        if not _port_free(self.port):
            raise TailorError(f"Port {self.port} is already in use; stop whatever uses it and retry.")
        key = config.api_key()
        # NO LLM_API_BASE: Resume-Matcher's native "openrouter" provider knows its endpoint; passing .../api/v1 makes
        # litellm call .../api/v1/v1/... and every request 404s (verified live).
        env = {k: v for k, v in os.environ.items() if k != "LLM_API_BASE"}
        env.update(LLM_PROVIDER="openrouter", LLM_MODEL=self.model, LLM_API_KEY=key, PYTHONUNBUFFERED="1",
                   FRONTEND_BASE_URL=f"http://127.0.0.1:{self.front_port}")
        log = open(config.DATA_DIR / "resume-matcher.log", "ab") if config.DATA_DIR.exists() else subprocess.DEVNULL
        self.proc = subprocess.Popen([str(RM_BACKEND / ".venv" / "bin" / "python"), "-m", "uvicorn", "app.main:app", "--port", str(self.port)],
                                     cwd=RM_BACKEND, env=env, stdout=log, stderr=log)
        t0 = time.time()
        while time.time() - t0 < wait:
            if self.proc.poll() is not None:
                raise TailorError("Resume-Matcher backend exited at startup (see brain/data/resume-matcher.log)")
            try:
                if requests.get(f"{self.base}/health", timeout=2, proxies={"http": None, "https": None}).status_code == 200:
                    return self
            except requests.RequestException:
                time.sleep(0.5)
        self.stop()
        raise TailorError("Resume-Matcher backend did not become ready in time")

    def start_frontend(self, wait: int = 180) -> None:
        """PDF export renders Resume-Matcher's /print page, so its Next.js frontend must be up (started on demand)."""
        if self.front:
            return
        node = nodeenv.node_bin()
        fe = RM_BACKEND.parent / "frontend"
        if not node or not (fe / "node_modules").exists():
            raise TailorError("Resume-Matcher frontend is not installed (apps/frontend/node_modules) or Node 22 is missing.")
        if not _port_free(self.front_port):
            raise TailorError(f"Port {self.front_port} is already in use; stop whatever uses it and retry.")
        env = {k: v for k, v in os.environ.items() if k not in ("PORT",)}
        env.update(NEXT_PUBLIC_API_URL=f"http://127.0.0.1:{RM_PORT}", BACKEND_ORIGIN=f"http://127.0.0.1:{RM_PORT}",
                   NEXT_TELEMETRY_DISABLED="1", PATH=str(Path(node).parent) + os.pathsep + env.get("PATH", ""))
        nxt = str(fe / "node_modules" / "next" / "dist" / "bin" / "next")
        log = open(config.DATA_DIR / "resume-matcher.log", "ab") if config.DATA_DIR.exists() else subprocess.DEVNULL
        if not (fe / ".next" / "BUILD_ID").exists():          # one-time production build (needs the self-hosted fonts patch)
            b = subprocess.run([node, nxt, "build"], cwd=fe, env=env, stdout=log, stderr=log, timeout=900)
            if b.returncode != 0:
                raise TailorError("Resume-Matcher frontend build failed (see brain/data/resume-matcher.log)")
        self.front = subprocess.Popen([node, nxt, "start", "-p", str(self.front_port), "-H", "127.0.0.1"],
                                      cwd=fe, env=env, stdout=log, stderr=log, start_new_session=True)
        t0 = time.time()
        while time.time() - t0 < wait:
            if self.front.poll() is not None:
                raise TailorError("Resume-Matcher frontend exited at startup (see brain/data/resume-matcher.log)")
            try:
                requests.get(f"http://127.0.0.1:{self.front_port}/", timeout=60, proxies={"http": None, "https": None})
                return
            except requests.RequestException:
                time.sleep(1)
        raise TailorError("Resume-Matcher frontend did not become ready in time")

    def stop(self) -> None:
        if self.front and self.front.poll() is None:
            import signal
            try:
                os.killpg(self.front.pid, signal.SIGTERM)
                self.front.wait(10)
            except (subprocess.TimeoutExpired, ProcessLookupError):
                try:
                    os.killpg(self.front.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
        self.front = None
        if self.proc and self.proc.poll() is None:
            self.proc.terminate()
            try:
                self.proc.wait(10)
            except subprocess.TimeoutExpired:
                self.proc.kill()
        self.proc = None

    def __enter__(self):
        return self.start()

    def __exit__(self, *a):
        self.stop()

    def req(self, method: str, path: str, timeout: int = 120, **kw):
        r = requests.request(method, self.base + path, timeout=timeout, proxies={"http": None, "https": None}, **kw)
        if r.status_code >= 400:
            try:
                detail = r.json().get("detail", r.text)
            except ValueError:
                detail = r.text
            raise TailorError(f"Resume-Matcher {method} {path} -> HTTP {r.status_code}: {str(detail)[:200]}")
        return r


def _rm_state_path(name: str) -> Path:
    return profiles.workspace_path(name, "data") / "resume_matcher.json"


def _has_content(data: dict) -> bool:
    """A parse that 'succeeded' with no name and no experience/education/projects is a failed parse (seen live: blank PDF)."""
    pr = (data or {}).get("processed_resume") or {}
    if not (pr.get("personalInfo") or {}).get("name"):
        return False
    return bool(pr.get("workExperience") or pr.get("education") or pr.get("personalProjects"))


def _ensure_master(api: Backend, name: str, cv_text: str) -> str:
    """Upload the profile's CV once (parsing costs LLM calls) and reuse it while cv.md is unchanged."""
    digest_ = hashlib.sha1(cv_text.encode()).hexdigest()
    sp = _rm_state_path(name)
    try:
        st = json.loads(sp.read_text())
    except (FileNotFoundError, json.JSONDecodeError):
        st = {}
    if st.get("cv_hash") == digest_ and st.get("resume_id"):
        try:
            d = api.req("GET", "/resumes", params={"resume_id": st["resume_id"]}).json()["data"]
            if d["raw_resume"]["processing_status"] == "ready" and _has_content(d):
                return st["resume_id"]
        except TailorError:
            pass                                                   # backend data was wiped -> upload again
    docx_path = cv_to_docx(cv_text, config.DATA_DIR / f"cv-{name}.docx")
    with open(docx_path, "rb") as f:
        up = api.req("POST", "/resumes/upload", files={"file": (f"{name}-cv.docx", f, DOCX_MIME)}, timeout=120).json()
    rid = up["resume_id"]
    t0 = time.time()
    while time.time() - t0 < 300:
        d = api.req("GET", "/resumes", params={"resume_id": rid}).json()["data"]
        status = d["raw_resume"]["processing_status"]
        if status == "ready":
            if not _has_content(d):
                raise TailorError("Resume-Matcher parsed your CV into an empty resume (this model cannot structure it)")
            sp.write_text(json.dumps({"cv_hash": digest_, "resume_id": rid}))
            return rid
        if status == "failed":
            raise TailorError("Resume-Matcher could not parse your CV (the model's answer was unusable)")
        time.sleep(3)
    raise TailorError("Resume-Matcher took too long to parse your CV")


def _rm_candidates(models: list[str] | None) -> list[str]:
    if models:
        return models
    pref = [m.strip() for m in (config.get("RESUME_MATCHER_MODELS") or "xiaomi/mimo-v2.6-flash").split(",") if m.strip()]
    rest = [m for m in config.models() if not m.endswith(":batch") and m not in pref]
    return pref + rest


def make_resume_matcher(job_id: str, name: str | None = None, models: list[str] | None = None, max_models: int = 3) -> dict:
    name = name or profiles.get_active()
    res = {"ok": False, "engine": "resume-matcher", "pdf_path": None, "error": None, "model": None}
    errors = {}
    try:
        job = _job(job_id, name)
        cv_text = profiles.read_file(name, "cv.md")
        fmt = "LETTER" if paper_format(name) == "letter" else "A4"
        for model in _rm_candidates(models)[:max_models]:
            try:
                with Backend(model) as api:
                    chk = api.req("POST", "/config/llm-test", timeout=90).json()
                    if not chk.get("healthy", False):
                        raise TailorError(f"model check failed: {str(chk)[:120]}")
                    rid = _ensure_master(api, name, cv_text)
                    jd = f"{job.get('title') or ''} at {job.get('company') or ''}\n\n{job['description']}"
                    jid = api.req("POST", "/jobs/upload", json={"job_descriptions": [jd], "resume_id": rid}).json()["job_id"][0]
                    imp = api.req("POST", "/resumes/improve", json={"resume_id": rid, "job_id": jid}, timeout=600).json()
                    new_id = imp["data"]["resume_id"]
                    if not _has_content(api.req("GET", "/resumes", params={"resume_id": new_id}).json()["data"]):
                        raise TailorError("Resume-Matcher produced an empty tailored resume")
                    api.start_frontend()
                    pdf = api.req("GET", f"/resumes/{new_id}/pdf", params={"pageSize": fmt}, timeout=180)
                    if not pdf.content.startswith(b"%PDF"):
                        raise TailorError("Resume-Matcher returned something that is not a PDF")
                    d = out_dir(job_id, name)
                    target = d / "resume-matcher.pdf"
                    _archive(target)
                    target.write_bytes(pdf.content)
                    res.update(ok=True, pdf_path=str(target), model=model)
                    return res
            except (TailorError, requests.RequestException) as e:
                errors[model] = str(e)[:160].replace(config.get("OPENROUTER_API_KEY") or "\0", "***")
        raise TailorError("Resume-Matcher failed with every model tried: " + "; ".join(f"{k}: {v}" for k, v in errors.items()))
    except TailorError as e:
        res["error"] = str(e)
    return res


# ================================================================== public entry point
def make_resume(job_id: str, engine: str = "career-ops", name: str | None = None) -> dict:
    name = name or profiles.get_active()
    if engine not in (*ENGINES, "both"):
        raise ValueError(f"engine must be one of {[*ENGINES, 'both']}")
    from brain import budget
    try:
        budget.require()
    except budget.BudgetExceeded as e:
        return {"ok": False, "engine": engine, "pdf_path": None, "model": None, "error": str(e)}
    if engine == "career-ops":
        return make_career_ops(job_id, name)
    if engine == "resume-matcher":
        return make_resume_matcher(job_id, name)
    results = {e: (make_career_ops if e == "career-ops" else make_resume_matcher)(job_id, name) for e in ENGINES}
    ok = [r for r in results.values() if r["ok"]]
    return {"ok": bool(ok), "engine": "both", "results": results, "pdf_path": ok[0]["pdf_path"] if ok else None,
            "pdf_paths": [r["pdf_path"] for r in ok], "model": None,
            "error": None if len(ok) == 2 else "; ".join(f"{e}: {r['error']}" for e, r in results.items() if not r["ok"])}
