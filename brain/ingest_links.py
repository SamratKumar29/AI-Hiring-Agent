"""Find job links in files/text and classify them.

extract_urls(source)  -> ordered, de-duplicated list of URLs (from .txt/.md/.csv/.html/.json text, .pdf text AND
                         clickable link annotations, .docx text AND hyperlinks, or a pasted string)
classify(url)         -> {url, type, board, ats, company}
    type: single_posting | board_search | company_careers | unknown
extract_and_classify(source) -> list of classify() dicts (preview before saving)
"""
import re
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse

URL_RE = re.compile(r"https?://[^\s<>\"'`\)\]\}\\]+", re.I)
WWW_RE = re.compile(r"(?<![\w/@.])(www\.[a-z0-9-]+(?:\.[a-z0-9-]+)+(?:/[^\s<>\"'`\)\]\}]*)?)", re.I)
TRACKING_PARAMS = {"gclid", "fbclid", "trk", "trkinfo", "trackingid", "refid", "lipi", "ref", "mc_cid",
                   "mc_eid", "src", "campaign", "ved", "usg", "utm_source", "utm_medium", "utm_campaign",
                   "utm_term", "utm_content", "utm_id", "source_ref", "from"}
TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".html", ".htm", ".json", ".rtf", ".url", ""}

# host suffix/substring -> board name (for job *boards*, as opposed to a company's own ATS)
BOARDS = [("linkedin.com", "linkedin"), ("indeed.", "indeed"), ("naukri.com", "naukri"),
          ("wellfound.com", "wellfound"), ("angel.co", "wellfound"), ("remoteok.", "remoteok"),
          ("michaelpage.", "michaelpage"), ("multirecruit", "multirecruit"), ("weworkremotely.com", "weworkremotely"),
          ("remotive.", "remotive"), ("glassdoor.", "glassdoor"), ("foundit.in", "foundit"),
          ("instahyre.com", "instahyre"), ("cutshort.io", "cutshort"), ("hirist.", "hirist"),
          ("workatastartup.com", "workatastartup")]

JOB_QUERY_KEYS = {"jobid", "job_id", "jid", "gh_jid", "jk", "currentjobid", "lever-source", "requisitionid", "reqid"}
SEARCH_WORDS = {"search", "results", "category", "categories", "all", "list", "browse", "filter"}


# ---------------------------------------------------------------- extraction
def _clean(u: str) -> str:
    return u.strip().rstrip(".,;:!?'\"*_").rstrip(")")


def _urls_from_text(text: str) -> list[str]:
    found = [(m.start(), _clean(m.group(0))) for m in URL_RE.finditer(text)]
    found += [(m.start(), "https://" + _clean(m.group(1))) for m in WWW_RE.finditer(text) if "://" not in m.group(1)]
    return [u for _, u in sorted(found) if len(u) > 10]       # keep the order they appear in


def _repair_wrapped_urls(text: str) -> str:
    """PDF text wraps long URLs across lines; re-join when a URL line ends mid-URL."""
    lines, out, i = text.splitlines(), [], 0
    while i < len(lines):
        cur = lines[i]
        while (re.search(r"https?://\S*$", cur) and re.search(r"[/\-_=&?%.]$", cur) and i + 1 < len(lines)
               and lines[i + 1].strip() and " " not in lines[i + 1].strip()):
            i += 1
            cur += lines[i].strip()
        out.append(cur)
        i += 1
    return "\n".join(out)


def _read_pdf(path: Path) -> list[str]:
    from pypdf import PdfReader

    urls: list[str] = []
    reader = PdfReader(str(path))
    for page in reader.pages:
        try:
            for annot in page.get("/Annots") or []:
                a = annot.get_object().get("/A")
                uri = a.get_object().get("/URI") if a is not None else None
                if uri:
                    urls.append(str(uri))
        except Exception:  # noqa: BLE001
            pass
        urls += _urls_from_text(_repair_wrapped_urls(page.extract_text() or ""))
    return urls


def _read_docx(path: Path) -> list[str]:
    import docx

    d = docx.Document(str(path))
    urls = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            text += "\n" + " ".join(c.text for c in row.cells)
    return urls + _urls_from_text(text)


def normalize_url(url: str) -> str:
    """Canonical form used for de-duplication: lower-case host, no fragment/tracking params/trailing slash."""
    p = urlparse(url.strip())
    host = p.netloc.lower()
    q = [(k, v) for k, v in parse_qs(p.query, keep_blank_values=False).items() if k.lower() not in TRACKING_PARAMS]
    query = urlencode([(k, vv) for k, vals in q for vv in vals])
    path = p.path.rstrip("/")
    return urlunparse((p.scheme.lower() or "https", host, path, "", query, ""))


def extract_urls(source) -> list[str]:
    """`source` = path to a file, or a pasted string. Returns de-duplicated URLs in first-seen order."""
    raw: list[str]
    p = None
    if not isinstance(source, str) or "\n" not in source and len(source) < 1024:
        try:
            cand = Path(source).expanduser()
            p = cand if cand.is_file() else None
        except (OSError, ValueError):
            p = None
    if p is not None:
        ext = p.suffix.lower()
        if ext == ".pdf":
            raw = _read_pdf(p)
        elif ext == ".docx":
            raw = _read_docx(p)
        elif ext in TEXT_EXT:
            raw = _urls_from_text(p.read_text(encoding="utf-8", errors="replace"))
        else:
            raise ValueError(f"Unsupported file type '{ext}'. Use .txt .md .csv .html .pdf or .docx")
    elif isinstance(source, str) and re.match(r"^(\S+\.)?(pdf|docx|txt|md)$", source.strip(), re.I):
        raise FileNotFoundError(f"File not found: {source}")
    else:
        raw = _urls_from_text(str(source))
    seen, out = set(), []
    for u in raw:
        if not re.match(r"https?://[^/\s]+\.[^/\s]+", u):
            continue
        key = normalize_url(u)
        if key not in seen:
            seen.add(key)
            out.append(u.strip())
    return out


# ---------------------------------------------------------------- classification
def slug_to_name(slug: str) -> str:
    return re.sub(r"[-_]+", " ", slug).strip().title() if slug else ""


def _result(url, typ, board=None, ats=None, company=None, canonical=None):
    return {"url": canonical or normalize_url(url), "type": typ, "board": board, "ats": ats,
            "company": slug_to_name(company) if company else None}


def _segments(path: str) -> list[str]:
    return [s for s in path.split("/") if s]


def _ats(url, host, segs, q):
    """Company-owned applicant tracking systems career-ops can scan. Returns result or None."""
    h = host
    if h.endswith("greenhouse.io") and (h.startswith(("boards.", "job-boards.", "boards-api."))):
        co = segs[0] if segs else None
        return _result(url, "single_posting" if len(segs) >= 3 and segs[1] == "jobs" else "company_careers",
                       ats="greenhouse", company=co)
    if h in ("jobs.lever.co", "jobs.eu.lever.co"):
        return _result(url, "single_posting" if len(segs) >= 2 else "company_careers", ats="lever",
                       company=segs[0] if segs else None)
    if h == "jobs.ashbyhq.com":
        return _result(url, "single_posting" if len(segs) >= 2 else "company_careers", ats="ashby",
                       company=segs[0] if segs else None)
    if h == "apply.workable.com":
        return _result(url, "single_posting" if len(segs) >= 3 and segs[1] == "j" else "company_careers",
                       ats="workable", company=segs[0] if segs else None)
    if h.endswith(".recruitee.com"):
        return _result(url, "single_posting" if segs[:1] == ["o"] and len(segs) > 1 else "company_careers",
                       ats="recruitee", company=h.split(".")[0])
    if h.endswith(".myworkdayjobs.com"):
        posting = "job" in segs or "details" in segs
        return _result(url, "single_posting" if posting else "company_careers", ats="workday", company=h.split(".")[0])
    if h in ("careers.smartrecruiters.com", "jobs.smartrecruiters.com"):
        return _result(url, "single_posting" if len(segs) >= 2 else "company_careers", ats="smartrecruiters",
                       company=segs[0] if segs else None)
    if h.endswith(".teamtailor.com"):
        return _result(url, "single_posting" if len(segs) >= 2 and segs[0] == "jobs" else "company_careers",
                       ats="teamtailor", company=h.split(".")[0])
    if h.endswith(".breezy.hr"):
        return _result(url, "single_posting" if segs[:1] == ["p"] else "company_careers", ats="breezy",
                       company=h.split(".")[0])
    if h.endswith(".bamboohr.com"):
        return _result(url, "single_posting" if len(segs) >= 2 and segs[0] == "careers" else "company_careers",
                       ats="bamboohr", company=h.split(".")[0])
    if h.endswith(".jobs.personio.de") or h.endswith(".jobs.personio.com"):
        return _result(url, "single_posting" if segs[:1] == ["job"] else "company_careers", ats="personio",
                       company=h.split(".")[0])
    if h == "jobs.jobvite.com":
        return _result(url, "single_posting" if "job" in segs else "company_careers", ats="jobvite",
                       company=segs[0] if segs else None)
    if h.endswith(".icims.com"):
        return _result(url, "single_posting" if "jobs" in segs and len(segs) > 2 else "company_careers",
                       ats="icims", company=h.split(".")[0].replace("careers-", ""))
    if "gh_jid" in q and not h.endswith("greenhouse.io"):
        return _result(url, "single_posting", ats="greenhouse")
    return None


def _board(url, host, segs, q, path):
    """Public job boards. Returns result or None."""
    board = next((b for suffix, b in BOARDS if suffix in host), None)
    if not board:
        return None
    post = lambda c=None, can=None: _result(url, "single_posting", board=board, company=c, canonical=can)  # noqa: E731
    search = lambda: _result(url, "board_search", board=board)  # noqa: E731
    if board == "linkedin":
        m = re.search(r"/jobs/view/(?:[^/]*-)?(\d{6,})", path)
        cur = (q.get("currentJobId") or q.get("currentjobid") or [None])[0]
        if m or cur:
            jid = m.group(1) if m else cur
            return post(can=f"https://www.linkedin.com/jobs/view/{jid}")
        return search() if "/jobs" in path else _result(url, "unknown", board=board)
    if board == "indeed":
        jk = (q.get("jk") or [None])[0]
        if ("/viewjob" in path or "/rc/clk" in path) and jk:
            return post(can=f"https://{host}/viewjob?jk={jk}")
        return search()
    if board == "naukri":
        return post() if re.search(r"/job-listings-.*-\d{6,}", path) else search()
    if board == "wellfound":
        return post() if re.search(r"/jobs/\d+", path) else search()
    if board == "remoteok":
        return post() if re.search(r"/remote-jobs/.*\d{5,}|/l/\d+", path) else search()
    if board == "michaelpage":
        return post() if re.search(r"/job-detail/|/job/|/ref[/_]", path) else search()
    if board == "weworkremotely":
        return post() if "/remote-jobs/" in path and len(segs) >= 2 else search()
    if board == "remotive":
        return post() if re.search(r"/remote-jobs/[^/]+/[^/]+", path) else search()
    if board == "glassdoor":
        return post() if "job-listing" in path or "jl=" in url else search()
    return post() if re.search(r"/jobs?/[^/]*\d", path) else search()


def classify(url: str) -> dict:
    p = urlparse(url.strip())
    host, path = p.netloc.lower().removeprefix("www."), p.path
    q = {k.lower(): v for k, v in parse_qs(p.query).items()}
    segs = _segments(path)
    for fn in (lambda: _ats(url, host, segs, q), lambda: _board(url, host, segs, q, path)):
        r = fn()
        if r:
            return r
    # generic company site
    if any(k in q for k in JOB_QUERY_KEYS):
        return _result(url, "single_posting")
    low = [s.lower() for s in segs]
    for i, seg in enumerate(low[:-1]):
        if seg in ("job", "jobs", "position", "positions", "opening", "openings", "vacancy", "vacancies",
                   "requisition", "posting", "postings", "careers", "career") and low[i + 1] not in SEARCH_WORDS:
            tail = segs[i + 1]
            if re.search(r"\d", tail) or tail.count("-") >= 2 or seg in ("job", "position", "opening", "vacancy"):
                return _result(url, "single_posting")
    if low and re.fullmatch(r"(careers?|jobs?|join-us|work-with-us|join|opportunities|vacancies|openings)", low[-1]):
        return _result(url, "company_careers", company=host.split(".")[-2] if "." in host else None)
    if low[:2] in (["about", "careers"], ["en", "careers"], ["company", "careers"]):
        return _result(url, "company_careers", company=host.split(".")[-2])
    return _result(url, "unknown")


def extract_and_classify(source) -> list[dict]:
    return [{**classify(u), "original": u} for u in extract_urls(source)]
