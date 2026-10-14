import json
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from types import SimpleNamespace

import pytest

from brain import digest, jobs_db, profiles, sync_env, tailor
from brain.llm import LLMUnavailable
from tests.report_samples import report

GOOD_HTML = "<!DOCTYPE html><html><body>" + "<p>Built production RAG systems with LangGraph and Qdrant for enterprise clients.</p>" * 30 + "</body></html>"


def test_validate_html():
    assert tailor.validate_html(GOOD_HTML) is None
    assert "does not start" in tailor.validate_html("Here is your CV: " + GOOD_HTML)
    assert "cut off" in tailor.validate_html(GOOD_HTML.replace("</html>", ""))
    assert "placeholders" in tailor.validate_html(GOOD_HTML.replace("Built", "{{NAME}}"))
    assert "tool-call" in tailor.validate_html(GOOD_HTML.replace("<p>", "<function=x><p>", 1))
    assert "almost no text" in tailor.validate_html("<!DOCTYPE html><html><body><p>hi</p></body></html>")


def test_cv_to_docx_keeps_structure(tmp_path):
    import docx
    p = tailor.cv_to_docx("# Jane Doe\n\n## Experience\n- Built **RAG** systems\n- Led a team\n\nPlain paragraph", tmp_path / "cv.docx")
    d = docx.Document(str(p))
    texts = [(x.style.name, x.text) for x in d.paragraphs]
    assert ("Heading 1", "Jane Doe") in texts and ("Heading 2", "Experience") in texts
    assert ("List Bullet", "Built RAG systems") in texts and ("Normal", "Plain paragraph") in texts


def test_paper_format_follows_country(active):
    assert tailor.paper_format(active) == "a4"                                  # fixture profile is in India
    profiles.update_preferences(active, {"country": "USA"})
    assert tailor.paper_format(active) == "letter"


@pytest.fixture
def job(active, monkeypatch, tmp_path):
    monkeypatch.setattr("brain.rank._fulltext", lambda url: None)
    jobs_db.upsert_jobs([{"url": "https://x.com/jobs/800001", "title": "AI Engineer", "company": "Acme",
                          "description": "RAG LangGraph Python " * 40, "source": "t"}], name=active)
    jid = jobs_db.job_id("https://x.com/jobs/800001")
    f = profiles.workspace_path(active, "reports") / "004-acme-2026-10-13.md"
    f.write_text(report(score="4.4"), encoding="utf-8")
    digest.rank_job(jid, f, active)
    return jid


def test_needs_a_report_and_a_description(active, monkeypatch):
    monkeypatch.setattr("brain.rank._fulltext", lambda url: None)
    jobs_db.upsert_jobs([{"url": "https://x.com/jobs/800002", "title": "T", "description": "RAG " * 50}], name=active)
    jid = jobs_db.job_id("https://x.com/jobs/800002")
    r = tailor.make_career_ops(jid, active)
    assert not r["ok"] and "rank" in r["error"]
    jobs_db.upsert_jobs([{"url": "https://x.com/jobs/800003", "title": "T"}], name=active)
    r = tailor.make_career_ops(jobs_db.job_id("https://x.com/jobs/800003"), active)
    assert not r["ok"] and "no usable description" in r["error"]
    assert not tailor.make_career_ops("missing", active)["ok"]
    with pytest.raises(ValueError):
        tailor.make_resume("x", "word")


@pytest.fixture
def fake_node(monkeypatch, tmp_path, active):
    """Fake the tailor script (writes HTML) and the PDF renderer; scripted per model."""
    co = tmp_path / "scratch-co"
    (co / "output").mkdir(parents=True)
    monkeypatch.setattr(tailor, "CAREER_OPS", co)
    monkeypatch.setattr(tailor.nodeenv, "node_bin", lambda: "node")
    state = SimpleNamespace(html=[GOOD_HTML], tried=[], render_ok=True, args=None)

    def fallback(cmd, cwd, timeout=900, models=None, validate=None):
        errors = {}
        for i, h in enumerate(state.html):
            m = f"model{i}"
            state.tried.append(m)
            if h is None:
                errors[m] = "HTTP 429"
                continue
            (co / "output" / "cv-test-acme.html").write_text(h, encoding="utf-8")
            import os, time
            os.utime(co / "output" / "cv-test-acme.html", (time.time() + i + 1, time.time() + i + 1))
            bad = validate(SimpleNamespace(returncode=0), m)
            if bad:
                errors[m] = bad
                continue
            return SimpleNamespace(returncode=0), m
        raise LLMUnavailable(errors)
    monkeypatch.setattr(sync_env, "run_with_fallback", fallback)

    def run_node(args, timeout=300):
        state.args = args
        if state.render_ok:
            Path(args[2]).write_bytes(b"%PDF-1.4 fake")
        return SimpleNamespace(returncode=0 if state.render_ok else 1, stdout="", stderr="" if state.render_ok else "Playwright: browser missing")
    monkeypatch.setattr(tailor, "_run_node", run_node)
    return state


def test_career_ops_engine_happy_path(active, job, fake_node):
    r = tailor.make_career_ops(job, active)
    assert r["ok"] and r["engine"] == "career-ops" and r["pdf_path"].endswith(f"/resumes/{job}/career-ops.pdf")
    assert Path(r["pdf_path"]).read_bytes().startswith(b"%PDF") and Path(r["html_path"]).exists()
    assert fake_node.args[0] == "generate-pdf.mjs" and "--format=a4" in fake_node.args and "--report=004" in fake_node.args
    assert tailor.list_resumes(job, active) == {"career-ops": r["pdf_path"]}


def test_unusable_html_falls_to_next_model_and_junk_is_removed(active, job, fake_node):
    fake_node.html = ["Sorry, I cannot do that.", GOOD_HTML]
    r = tailor.make_career_ops(job, active)
    assert r["ok"] and r["model"] == "model1" and fake_node.tried == ["model0", "model1"]


def test_all_models_unusable_gives_a_clear_error(active, job, fake_node):
    fake_node.html = ["nope", None]
    r = tailor.make_career_ops(job, active)
    assert not r["ok"] and "No model produced a usable CV" in r["error"] and r["pdf_path"] is None


def test_render_failure_is_reported_and_leaves_no_pdf(active, job, fake_node):
    fake_node.render_ok = False
    r = tailor.make_career_ops(job, active)
    assert not r["ok"] and "PDF rendering failed" in r["error"] and "browser missing" in r["error"]
    assert tailor.list_resumes(job, active) == {}


def test_regenerating_archives_the_previous_pdf(active, job, fake_node):
    first = tailor.make_career_ops(job, active)["pdf_path"]
    Path(first).write_bytes(b"%PDF-1.4 OLD")
    tailor.make_career_ops(job, active)
    hist = list((Path(first).parent / "history").glob("career-ops-*.pdf"))
    assert len(hist) == 1 and hist[0].read_bytes() == b"%PDF-1.4 OLD"


def test_both_engines_report_each_result(active, job, fake_node, monkeypatch):
    monkeypatch.setattr(tailor, "make_resume_matcher", lambda j, n: {"ok": False, "engine": "resume-matcher", "pdf_path": None, "error": "backend down", "model": None})
    r = tailor.make_resume(job, "both", active)
    assert r["ok"] and len(r["pdf_paths"]) == 1 and "resume-matcher: backend down" in r["error"]
    assert r["results"]["career-ops"]["ok"] and not r["results"]["resume-matcher"]["ok"]


# ---------------- Resume-Matcher engine against a fake local server ----------------
class FakeRM(BaseHTTPRequestHandler):
    log = []
    llm_ok = True

    def _send(self, code, body, ctype="application/json"):
        data = body if isinstance(body, bytes) else json.dumps(body).encode()
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, *a):
        pass

    def do_GET(self):
        FakeRM.log.append(("GET", self.path))
        if self.path.endswith("/health"):
            return self._send(200, {"status": "healthy"})
        if "/resumes?resume_id=" in self.path:
            return self._send(200, {"data": {"raw_resume": {"processing_status": "ready"}, "processed_resume": FakeRM.processed}})
        if self.path.startswith("/api/v1/resumes/improved-1/pdf"):
            return self._send(200, b"%PDF-1.7 tailored", "application/pdf")
        self._send(404, {"detail": "nope"})

    def do_POST(self):
        n = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(n)
        FakeRM.log.append(("POST", self.path, len(body)))
        if self.path.endswith("/config/llm-test"):
            return self._send(200, {"healthy": FakeRM.llm_ok})
        if self.path.endswith("/resumes/upload"):
            return self._send(200, {"resume_id": "master-1", "processing_status": "pending"})
        if self.path.endswith("/jobs/upload"):
            return self._send(200, {"job_id": ["job-1"]})
        if self.path.endswith("/resumes/improve"):
            return self._send(200, {"data": {"resume_id": "improved-1"}})
        self._send(404, {"detail": "nope"})


@pytest.fixture
def rm_server(monkeypatch):
    srv = HTTPServer(("127.0.0.1", 0), FakeRM)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    FakeRM.log, FakeRM.llm_ok = [], True
    FakeRM.processed = {"personalInfo": {"name": "A B"}, "workExperience": [{"title": "x"}]}
    monkeypatch.setattr(tailor.Backend, "start_frontend", lambda self, wait=180: None)
    port = srv.server_address[1]

    class FakeBackend(tailor.Backend):
        def __init__(self, model):
            super().__init__(model, port)
            self.models_used.append(model)

        def start(self, wait=60):
            return self

        def stop(self):
            pass
    FakeBackend.models_used = []
    monkeypatch.setattr(tailor, "Backend", FakeBackend)
    yield FakeBackend
    srv.shutdown()


def test_resume_matcher_flow_uploads_once_and_saves_pdf(active, job, rm_server, tmp_path, monkeypatch):
    monkeypatch.setattr(tailor.config, "DATA_DIR", tmp_path)
    r = tailor.make_resume_matcher(job, active, models=["m/one"])
    assert r["ok"] and r["model"] == "m/one" and Path(r["pdf_path"]).read_bytes().startswith(b"%PDF")
    posts = [e[1].split("/api/v1")[1] for e in FakeRM.log if e[0] == "POST"]
    assert posts == ["/config/llm-test", "/resumes/upload", "/jobs/upload", "/resumes/improve"]
    FakeRM.log.clear()
    tailor.make_resume_matcher(job, active, models=["m/one"])                      # CV unchanged -> no second upload
    assert "/resumes/upload" not in " ".join(str(e[1]) for e in FakeRM.log if e[0] == "POST")
    assert (profiles.workspace_path(active, "data") / "resume_matcher.json").exists()


def test_resume_matcher_reuploads_when_cv_changes(active, job, rm_server, tmp_path, monkeypatch):
    monkeypatch.setattr(tailor.config, "DATA_DIR", tmp_path)
    tailor.make_resume_matcher(job, active, models=["m/one"])
    profiles.write_file(active, "cv.md", "# New CV\nCompletely different content here.")
    FakeRM.log.clear()
    tailor.make_resume_matcher(job, active, models=["m/one"])
    assert any(e[0] == "POST" and e[1].endswith("/resumes/upload") for e in FakeRM.log)


def test_resume_matcher_tries_next_model_when_llm_check_fails(active, job, rm_server, tmp_path, monkeypatch):
    monkeypatch.setattr(tailor.config, "DATA_DIR", tmp_path)
    FakeRM.llm_ok = False
    r = tailor.make_resume_matcher(job, active, models=["m/a", "m/b"])
    assert not r["ok"] and "m/a: model check failed" in r["error"] and "m/b: model check failed" in r["error"]
    assert rm_server.models_used == ["m/a", "m/b"]


def test_backend_env_never_sets_api_base_and_drops_inherited_one(monkeypatch, tmp_path):
    """Live-verified: a LLM_API_BASE of .../api/v1 makes every OpenRouter call 404 inside Resume-Matcher."""
    seen = {}

    class FakePopen:
        def __init__(self, cmd, cwd=None, env=None, stdout=None, stderr=None):
            seen["env"] = env
        def poll(self):
            return None
        def terminate(self): pass
        def wait(self, t=None): pass
    monkeypatch.setattr(tailor.subprocess, "Popen", FakePopen)
    monkeypatch.setattr(tailor, "RM_BACKEND", tmp_path)
    (tmp_path / ".venv" / "bin").mkdir(parents=True)
    (tmp_path / ".venv" / "bin" / "python").write_text("")
    monkeypatch.setattr(tailor.config, "api_key", lambda: "sk-test")
    monkeypatch.setattr(tailor.requests, "get", lambda *a, **k: SimpleNamespace(status_code=200))
    monkeypatch.setenv("LLM_API_BASE", "https://openrouter.ai/api/v1")
    b = tailor.Backend("some/model", port=59123)
    b.start()
    assert "LLM_API_BASE" not in seen["env"] and seen["env"]["LLM_PROVIDER"] == "openrouter"
    assert seen["env"]["LLM_MODEL"] == "some/model" and seen["env"]["LLM_API_KEY"] == "sk-test"
    b.stop()


def test_resume_matcher_rejects_empty_parse(active, job, rm_server, tmp_path, monkeypatch):
    monkeypatch.setattr(tailor.config, "DATA_DIR", tmp_path)
    FakeRM.processed = {"personalInfo": {"name": ""}, "workExperience": []}
    r = tailor.make_resume_matcher(job, active, models=["m/one"])
    assert not r["ok"] and "empty" in r["error"]


def test_node_children_get_proxy_support(monkeypatch):
    from brain import nodeenv
    monkeypatch.setenv("HTTPS_PROXY", "http://127.0.0.1:3128")
    monkeypatch.delenv("NODE_USE_ENV_PROXY", raising=False)
    assert nodeenv.env_with_node().get("NODE_USE_ENV_PROXY") == "1"
    monkeypatch.delenv("HTTPS_PROXY")
    for v in nodeenv.PROXY_VARS:
        monkeypatch.delenv(v, raising=False)
    assert "NODE_USE_ENV_PROXY" not in nodeenv.env_with_node()
